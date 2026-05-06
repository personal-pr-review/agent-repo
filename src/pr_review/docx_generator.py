from __future__ import annotations

from pathlib import Path

from docx import Document

from .models import ComparisonResult, PRMetadata, ReviewResult


class DocxGenerator:
    def __init__(self, template_path: Path) -> None:
        self._template_path = template_path

    def generate(
        self,
        output_path: Path,
        pr_metadata: PRMetadata,
        comparisons: list[ComparisonResult],
        review_result: ReviewResult,
    ) -> None:
        if not self._template_path.exists():
            raise FileNotFoundError(f"Template not found: {self._template_path}")

        document = Document(str(self._template_path))
        if len(document.tables) < 4:
            raise ValueError("Template structure mismatch: expected at least 4 tables.")

        self._fill_document_header_table(document.tables[0], pr_metadata)
        self._fill_pr_metadata_table(document.tables[1], pr_metadata)
        self._fill_files_changed_table(document.tables[2], comparisons)
        self._fill_behavior_change_table(document.tables[3], comparisons)
        self._fill_narrative_sections(document, pr_metadata, comparisons, review_result)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))

    @staticmethod
    def _fill_document_header_table(table, pr_metadata: PRMetadata) -> None:
        values = [
            pr_metadata.repository,
            "",
            pr_metadata.title,
            f"{pr_metadata.pr_number} / {pr_metadata.html_url}",
            f"{pr_metadata.head_branch} \u2192 {pr_metadata.base_branch}",
            pr_metadata.author,
            "",
            "",
            "",
        ]
        for idx, value in enumerate(values, start=1):
            DocxGenerator._set_cell_text(table.cell(idx, 1), value)

    @staticmethod
    def _fill_pr_metadata_table(table, pr_metadata: PRMetadata) -> None:
        values = [
            str(pr_metadata.changed_files),
            str(pr_metadata.additions),
            str(pr_metadata.deletions),
            str(pr_metadata.commits),
            pr_metadata.linked_ticket,
            pr_metadata.release_sprint,
        ]
        for idx, value in enumerate(values, start=1):
            DocxGenerator._set_cell_text(table.cell(idx, 1), value)

    @staticmethod
    def _fill_files_changed_table(table, comparisons: list[ComparisonResult]) -> None:
        DocxGenerator._reset_data_rows(table)

        if not comparisons:
            row = table.add_row().cells
            DocxGenerator._set_cell_text(row[0], "")
            DocxGenerator._set_cell_text(row[1], "")
            DocxGenerator._set_cell_text(row[2], "")
            DocxGenerator._set_cell_text(row[3], "")
            return

        for item in comparisons:
            row = table.add_row().cells
            DocxGenerator._set_cell_text(row[0], item.file_path)
            DocxGenerator._set_cell_text(row[1], item.change_type)
            DocxGenerator._set_cell_text(row[2], item.change_summary)
            DocxGenerator._set_cell_text(row[3], item.risk_level)

    @staticmethod
    def _fill_behavior_change_table(table, comparisons: list[ComparisonResult]) -> None:
        DocxGenerator._reset_data_rows(table)

        if not comparisons:
            row = table.add_row().cells
            for idx in range(5):
                DocxGenerator._set_cell_text(row[idx], "")
            return

        for item in comparisons:
            row = table.add_row().cells
            DocxGenerator._set_cell_text(row[0], item.file_path)
            DocxGenerator._set_cell_text(row[1], "")
            DocxGenerator._set_cell_text(row[2], item.behavioral_change)
            DocxGenerator._set_cell_text(row[3], item.change_summary)
            DocxGenerator._set_cell_text(row[4], item.semantic_impact)

    @staticmethod
    def _reset_data_rows(table) -> None:
        while len(table.rows) > 1:
            table._tbl.remove(table.rows[-1]._tr)

    @staticmethod
    def _set_cell_text(cell, value: str) -> None:
        text = value or ""
        if cell.paragraphs and cell.paragraphs[0].runs:
            cell.paragraphs[0].runs[0].text = text
            for run in cell.paragraphs[0].runs[1:]:
                run.text = ""
            for paragraph in cell.paragraphs[1:]:
                paragraph.text = ""
            return
        cell.text = text

    @staticmethod
    def _fill_narrative_sections(
        document: Document,
        pr_metadata: PRMetadata,
        comparisons: list[ComparisonResult],
        review_result: ReviewResult,
    ) -> None:
        derived_summary = review_result.summary or "; ".join(
            [item.change_summary for item in comparisons if item.change_summary][:3]
        )
        purpose = review_result.purpose_of_pr or f"This PR updates {pr_metadata.title}."
        summary_of_changes = review_result.summary_of_changes or derived_summary
        problem = review_result.problem_being_solved or derived_summary
        expected = review_result.expected_outcome or "Expected behavior should align with intended PR outcomes."
        risk = review_result.risk_level or "Low"
        decision = review_result.final_recommendation or "Merge"
        reasoning = review_result.reasoning or derived_summary

        for paragraph in document.paragraphs:
            text = paragraph.text
            updated = text

            updated = updated.replace("[Explain why this PR was created]", purpose)
            updated = updated.replace("[Brief summary of key changes]", summary_of_changes)
            updated = updated.replace("[Low / Medium / High]", risk)
            updated = updated.replace("[Merge / Do Not Merge]", decision)

            if updated.strip() == "Problem Being Solved:":
                updated = f"Problem Being Solved: {problem}"
            elif updated.startswith("Problem Being Solved:") and updated.strip() == "Problem Being Solved:":
                updated = f"Problem Being Solved: {problem}"

            if updated.strip() == "Expected Outcome:":
                updated = f"Expected Outcome: {expected}"
            elif updated.startswith("Expected Outcome:") and updated.strip() == "Expected Outcome:":
                updated = f"Expected Outcome: {expected}"

            if "Decision:" in updated and "Reasoning:" in updated:
                updated = f"Decision: {decision}\nReasoning: {reasoning}"
            elif updated.strip() == "Reasoning:":
                updated = f"Reasoning: {reasoning}"
            elif updated.startswith("Reasoning:") and updated.strip() == "Reasoning:":
                updated = f"Reasoning: {reasoning}"

            if updated != text:
                DocxGenerator._set_paragraph_text(paragraph, updated)

    @staticmethod
    def _set_paragraph_text(paragraph, text: str) -> None:
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ""
            return
        paragraph.text = text
